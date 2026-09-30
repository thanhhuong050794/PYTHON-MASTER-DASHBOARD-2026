"""Xuất báo cáo tổng quan ra Excel, Word và PDF."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .filters import ContactFilters
from .security import CurrentUser


def _n(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "Có" if v else "Không"
    if isinstance(v, float):
        if abs(v) >= 100:
            return f"{v:,.0f}".replace(",", ".")
        return f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, int):
        return f"{v:,}".replace(",", ".")
    return str(v)


def _vnd(v) -> str:
    if v is None:
        return "—"
    return f"{int(round(v)):,} ₫".replace(",", ".")


def _pct(v) -> str:
    if v is None:
        return "—"
    sign = "+" if v > 0 else ""
    return f"{sign}{v * 100:,.1f}%".replace(",", "X").replace(".", ",").replace("X", ".")


def _filter_note(f: ContactFilters, viewing: str | None) -> str:
    parts = []
    if viewing:
        parts.append(f"Phạm vi đối tác: {viewing}")
    else:
        parts.append("Phạm vi: toàn hệ thống")
    if f.date_from or f.date_to:
        parts.append(f"Thời gian: {f.date_from or '…'} → {f.date_to or '…'}")
    if f.board:
        parts.append(f"Bảng: {f.board}")
    if f.region:
        parts.append(f"Miền: {f.region}")
    return " · ".join(parts)


def report_sections(data: dict, user: CurrentUser, f: ContactFilters) -> list[tuple[str, list[list[str]]]]:
    viewing = data.get("viewing_partner")
    generated = datetime.now().strftime("%d/%m/%Y %H:%M")
    sections: list[tuple[str, list[list[str]]]] = [
        ("Thông tin báo cáo", [
            ["Tiêu đề", "Báo cáo tổng quan tuyển sinh Python Master · PTIT"],
            ["Người xuất", f"{user.full_name} (@{user.username})"],
            ["Thời điểm xuất", generated],
            ["Bộ lọc", _filter_note(f, viewing)],
        ]),
    ]

    crm = data.get("crm")
    if crm:
        k = crm["kpi"]
        sections.append(("CRM — chỉ số", [
            ["Chỉ số", "Giá trị"],
            ["Liên hệ duy nhất", _n(k.get("total"))],
            ["Đăng ký trên website", _n(k.get("web_registered"))],
            ["Đã thanh toán", _n(k.get("paid"))],
            ["Tỉ lệ thanh toán / đăng ký", _pct(k.get("paid_rate"))],
            ["Doanh thu lệ phí ghi nhận", _vnd(k.get("revenue"))],
            ["Doanh thu tiềm năng", _vnd(k.get("potential_revenue"))],
            ["Chưa thanh toán (đã đăng ký)", _n(k.get("unpaid_registered"))],
        ]))
        sections.append(("Phễu tuyển sinh", [["Giai đoạn", "Số lượng"]] + [[r["stage"], _n(r["value"])] for r in crm.get("funnel", [])]))
        if crm.get("by_segment"):
            sections.append(("Segment CRM", [["Segment", "Liên hệ", "Đăng ký web", "Đã thanh toán", "Doanh thu"]] + [
                [f"{r.get('key') or '—'} · {r.get('name') or ''}".strip(" ·"), _n(r.get("count")), _n(r.get("web_registered")), _n(r.get("paid")), _vnd(r.get("revenue"))]
                for r in crm["by_segment"]
            ]))
        if crm.get("by_channel"):
            sections.append(("Kênh thu hút", [["Kênh", "Liên hệ", "Đăng ký web", "Đã thanh toán", "Doanh thu"]] + [
                [r.get("key") or "—", _n(r.get("count")), _n(r.get("web_registered")), _n(r.get("paid")), _vnd(r.get("revenue"))]
                for r in crm["by_channel"]
            ]))

    mk = data.get("marketing")
    if mk:
        k = mk["kpi"]
        sections.append(("Kênh & ROI — chỉ số", [
            ["Chỉ số", "Giá trị"],
            ["Ngân sách nhìn thấy", _vnd(k.get("budget"))],
            ["ROI toàn chương trình", _pct(k.get("roi_program"))],
            ["ROI kênh được giao", _pct(k.get("roi_paid_channels"))],
            ["CPA thí sinh đã đóng phí", _vnd(k.get("cpa_paid"))],
        ]))
        visible = [c for c in mk.get("channels", []) if c.get("budget_visible")]
        if visible:
            sections.append(("Ngân sách và doanh thu theo kênh", [["Kênh", "Ngân sách", "Doanh thu", "ROI", "CPA đóng phí", "Đã đóng phí"]] + [
                [c.get("channel") or "—", _vnd(c.get("budget")), _vnd(c.get("revenue")), _pct(c.get("roi")), _vnd(c.get("cpa_paid")), _n(c.get("paid"))]
                for c in visible
            ]))

    ex = data.get("exams")
    if ex:
        k = ex["kpi"]
        sections.append(("Kết quả thi — chỉ số", [
            ["Chỉ số", "Giá trị"],
            ["Thí sinh có điểm", _n(k.get("count"))],
            ["Điểm TB Vòng loại", _n(k.get("avg"))],
            ["Điểm cao nhất", _n(k.get("max"))],
            ["Vào chung kết", _n(k.get("finalists"))],
            ["Đạt chứng chỉ COS Pro", _n(k.get("certificates"))],
            ["Tỉ lệ chứng chỉ", _pct(k.get("cert_rate"))],
        ]))
        if ex.get("grades"):
            sections.append(("Xếp loại Vòng loại theo bảng", [["Bảng", "Xếp loại", "Thí sinh"]] + [
                [g.get("board") or "—", g.get("grade") or "—", _n(g.get("count"))] for g in ex["grades"]
            ]))

    sp = data.get("sponsorship")
    if sp:
        sections.append(("Nhà tài trợ & truyền thông", [
            ["Chỉ số", "Giá trị"],
            ["Tiến độ quyền lợi (TB)", _pct(sp.get("avg_progress"))],
            ["Tổng lượt tiếp cận", _n(sp.get("total_reach"))],
            ["Tổng lượt tương tác", _n(sp.get("total_engagement"))],
        ]))
        if sp.get("benefits"):
            sections.append(("Quyền lợi nhà tài trợ", [["Hạng mục", "Tiến độ", "Trạng thái"]] + [
                [b.get("item") or "—", _pct(b.get("progress")), b.get("status") or "—"] for b in sp["benefits"]
            ]))

    pipe = data.get("pipeline")
    if pipe:
        raw = pipe.get("raw_counts") or {}
        sections.append(("Dữ liệu thô → dữ liệu sạch", [
            ["Chỉ số", "Giá trị"],
            ["Dòng thô (3 nguồn)", _n(pipe.get("total_raw"))],
            ["Website", _n(raw.get("raw_web_signups"))],
            ["Talkshow", _n(raw.get("raw_talkshow"))],
            ["FB Lead", _n(raw.get("raw_fb_leads"))],
            ["Liên hệ duy nhất", _n(pipe.get("unique_contacts"))],
            ["Bản ghi bị gộp", _n(pipe.get("merged_records"))],
            ["Tỉ lệ trùng", _pct(pipe.get("duplicate_rate"))],
        ]))

    return sections


def _safe_sheet(name: str) -> str:
    bad = set(r"[]:*?/\\")
    cleaned = "".join("_" if c in bad else c for c in name)[:31]
    return cleaned or "Sheet"


def build_xlsx(sections: list[tuple[str, list[list[str]]]]) -> bytes:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1C5CAB")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=12, color="1C5CAB")
    thin = Border(
        left=Side(style="thin", color="C3C2B7"),
        right=Side(style="thin", color="C3C2B7"),
        top=Side(style="thin", color="C3C2B7"),
        bottom=Side(style="thin", color="C3C2B7"),
    )
    first = True
    for title, rows in sections:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = _safe_sheet(title)
        ws["A1"] = title
        ws["A1"].font = title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(rows[0]) if rows else 1, 1))
        for r_i, row in enumerate(rows, start=3):
            for c_i, val in enumerate(row, start=1):
                cell = ws.cell(r_i, c_i, val)
                cell.border = thin
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if r_i == 3:
                    cell.fill = header_fill
                    cell.font = header_font
        widths = [12] * max((len(r) for r in rows), default=1)
        for row in rows:
            for i, val in enumerate(row):
                widths[i] = min(42, max(widths[i], min(len(str(val)) + 2, 42)))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:3"
        ws.oddHeader.left.text = "NEU Admissions BI"
        ws.oddFooter.right.text = "Trang &P / &N"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_docx(sections: list[tuple[str, list[list[str]]]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    def _set_run_font(run, name="Calibri"):
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

    h = doc.add_heading("Báo cáo tổng quan tuyển sinh", level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
        _set_run_font(run)
    sub = doc.add_paragraph("Python Master · PTIT · NEU Admissions BI")
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in sub.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
        _set_run_font(run)

    for title, rows in sections:
        heading = doc.add_heading(title, level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
            _set_run_font(run)
        if not rows:
            continue
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        table.autofit = True
        for r_i, row in enumerate(rows):
            for c_i in range(cols):
                cell = table.cell(r_i, c_i)
                cell.text = row[c_i] if c_i < len(row) else ""
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        _set_run_font(run)
                        if r_i == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if r_i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:fill"), "1C5CAB")
                    shd.set(qn("w:val"), "clear")
                    cell._tc.get_or_add_tcPr().append(shd)
        doc.add_paragraph("")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_font() -> tuple[str, str | None]:
    candidates = [
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ]
    for p in candidates:
        if p.is_file():
            return "ReportVN", str(p)
    return "Helvetica", None


def build_pdf(sections: list[tuple[str, list[list[str]]]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, path = _pdf_font()
    if path:
        pdfmetrics.registerFont(TTFont(font, path))
        bold_path = path.replace("arial.ttf", "arialbd.ttf").replace("segoeui.ttf", "segoeuib.ttf").replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
        bold_name = font + "-Bold"
        if Path(bold_path).is_file():
            pdfmetrics.registerFont(TTFont(bold_name, bold_path))
        else:
            bold_name = font
    else:
        bold_name = "Helvetica-Bold"

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title="Báo cáo tổng quan tuyển sinh",
        author="NEU Admissions BI",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "VNTitle", parent=styles["Title"], fontName=bold_name, fontSize=16, textColor=colors.HexColor("#1C5CAB"),
        spaceAfter=4, leading=20,
    )
    h_style = ParagraphStyle(
        "VNHeading", parent=styles["Heading2"], fontName=bold_name, fontSize=11, textColor=colors.HexColor("#1C5CAB"),
        spaceBefore=10, spaceAfter=6, leading=14,
    )
    cell_style = ParagraphStyle("VNCell", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#0B0B0B"))
    cell_head = ParagraphStyle("VNHead", fontName=bold_name, fontSize=8, leading=11, textColor=colors.white)

    story = [
        Paragraph("Báo cáo tổng quan tuyển sinh — Python Master · PTIT", title_style),
        Paragraph("NEU Admissions BI", ParagraphStyle("sub", fontName=font, fontSize=9, textColor=colors.HexColor("#52514E"), spaceAfter=8)),
    ]
    navy = colors.HexColor("#1C5CAB")
    grid = colors.HexColor("#C3C2B7")
    zebra = colors.HexColor("#F3F2EE")

    for title, rows in sections:
        story.append(Paragraph(title, h_style))
        if not rows:
            continue
        data = []
        for r_i, row in enumerate(rows):
            style = cell_head if r_i == 0 else cell_style
            data.append([Paragraph(str(c).replace("&", "&amp;").replace("<", "&lt;"), style) for c in row])
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), bold_name),
            ("FONTNAME", (0, 1), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 1), (-1, -1), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, zebra]),
            ("GRID", (0, 0), (-1, -1), 0.4, grid),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 4))

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#7A7873"))
        canvas.drawString(12 * mm, 8 * mm, "NEU Admissions BI · Báo cáo tổng quan")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"Trang {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


_MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def report_response(kind: str, data: dict, user: CurrentUser, f: ContactFilters) -> StreamingResponse:
    sections = report_sections(data, user, f)
    builders = {"xlsx": build_xlsx, "docx": build_docx, "pdf": build_pdf}
    payload = builders[kind](sections)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"bao_cao_tong_quan_{stamp}.{kind}"
    return StreamingResponse(
        BytesIO(payload),
        media_type=_MIME[kind],
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
